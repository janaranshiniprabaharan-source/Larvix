from typing import Dict
import json

# PDF
try:  # Optional dependency
    from reportlab.lib.pagesizes import letter  # type: ignore
    from reportlab.pdfgen import canvas  # type: ignore
except ImportError:  # pragma: no cover
    letter = None  # type: ignore
    canvas = None  # type: ignore

# Word
try:  # Optional dependency
    from docx import Document  # type: ignore
except ImportError:  # pragma: no cover
    Document = None  # type: ignore

def json_report(report: Dict, output_file: str):
    with open(output_file, "w") as f:
        json.dump(report, f, indent=4)

def txt_report(report: Dict, output_file: str):
    lines = [
        "================ Larvix Report ================\n",
        f"File: {report['path']}",
        f"Size: {report['filesize']} bytes",
        f"MD5: {report['md5']}",
        f"SHA1: {report['sha1']}",
        f"SHA256: {report['sha256']}",
        f"Entropy: {report['entropy']}",
        f"Strings found: {report['num_strings']}",
        f"Embedded URLs: {report['urls']}",
        f"Embedded IPs: {report['ips']}",
        f"Score: {report['score']}",
        f"Verdict: {report['verdict']}",
        f"Reasons: {', '.join(report['reasons']) if report['reasons'] else 'None'}",
        f"Report time: {report['timestamp']}\n"
    ]
    with open(output_file, "w") as f:
        f.write("\n".join(lines))

def pdf_report(report: Dict, output_file: str):
    if canvas is None or letter is None:
        raise RuntimeError("PDF export requires the 'reportlab' package. Install with: pip install reportlab")
    c = canvas.Canvas(output_file, pagesize=letter)  # type: ignore
    width, height = letter  # type: ignore
    y = height - 40
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, y, "Larvix Report")
    y -= 30
    c.setFont("Helvetica", 12)
    for key, value in report.items():
        line = f"{key}: {value}"
        c.drawString(50, y, str(line)[:100])
        y -= 20
        if y < 50:
            c.showPage()
            y = height - 40
    c.save()

def docx_report(report: Dict, output_file: str):
    if Document is None:
        raise RuntimeError("DOCX export requires the 'python-docx' package. Install with: pip install python-docx")
    doc = Document()  # type: ignore
    doc.add_heading("Larvix Report", 0)
    for key, value in report.items():
        doc.add_paragraph(f"{key}: {value}")
    doc.save(output_file)
